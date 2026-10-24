"""Create implementation-specific revision supplements, preserving earlier documents."""
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,shutil,datetime
ROOT=Path(__file__).resolve().parents[1]
OUT=Path('D:/绩点/大三上/工业APP开发/第三四阶段交付');OUT.mkdir(exist_ok=True)
accept=json.loads((ROOT/'.runtime/stage34-acceptance.json').read_text(encoding='utf-8'))
remote=json.loads((ROOT/'.runtime/stage34-remote-acceptance.json').read_text(encoding='utf-8'))
fault=json.loads((ROOT/'.runtime/stage34-fault-acceptance.json').read_text(encoding='utf-8'))
def doc(title,subtitle):
 d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(2.1);s.bottom_margin=Cm(2);s.left_margin=Cm(2.2);s.right_margin=Cm(2.2)
 for name in ['Normal','Title','Heading 1','Heading 2']:
  st=d.styles[name];st.font.name='宋体' if name=='Normal' else '黑体';st._element.rPr.rFonts.set(qn('w:eastAsia'),st.font.name);st.font.color.rgb=RGBColor(0,0,0)
  st.font.size=Pt({'Normal':11,'Title':20,'Heading 1':15,'Heading 2':12}[name]);st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.15
 for st in d.styles:
  for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 d.add_paragraph(title,'Title');d.add_paragraph(subtitle)
 p=s.footer.paragraphs[0];p.alignment=2;r=p.add_run('MouseMES  第三四阶段修订  ');r.font.size=Pt(9)
 f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
 return d
def h(d,title):d.add_heading(title,1)
def p(d,text):d.add_paragraph(text)
def page(d,title):d.add_page_break();h(d,title)
def table(d,headers,rows,widths):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for c,text in zip(t.rows[0].cells,headers):c.text=text
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for row in rows:
  for c,text in zip(t.add_row().cells,row):c.text=str(text)
 for i,row in enumerate(t.rows):
  for j,c in enumerate(row.cells):
   c.width=Cm(widths[j]);pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b)
   if i==0:
    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'DCE6F1');pr.append(fill)
   for para in c.paragraphs:
    para.paragraph_format.space_before=Pt(4);para.paragraph_format.space_after=Pt(4)
    for r in para.runs:r.font.size=Pt(10);r.bold=i==0
 d.add_paragraph()
 return t

d=doc('MouseMES软件需求规格补充说明','第三阶段与第四阶段已实现需求  版本 3.4  2026年10月23日')
p(d,'本补充说明与原《MouseMES软件需求规格说明书》共同使用，修订原2.12节智能查询与2.13节定制功能的实现范围。本文列明本次已交付行为；原文中的设备采集、线边转移和分批发货等扩展目标不能据此宣称已经实现。')
h(d,'一 业务目标和使用角色')
p(d,'面向鼠标制造业务，用户以自然语言查询真实库存、当前缺料和工单追溯，并用相同条件的人工入口复核。生产报工提交后，系统自动创建质检待办，质检员从待办进入检测，合格后推进工序与工单进度。')
table(d,['能力','实际行为和验收条件'],[
 ('REQ AI 01 查询','支持库存、缺料、工单追溯三种白名单工具。模型或本地解析器只生成条件，不执行模型生成的SQL，不负责计算库存。'),
 ('REQ AI 02 准确性','库存按物料和仓库展示，逐行带单位；日期含首尾。展示当前、期初、入库、出库和期末；不能核实的余额明确留空并说明原因。'),
 ('REQ AI 03 多轮','保留会话物料、产品、日期、仓库等条件；“改查上个月”覆盖时间。多匹配仓库或物料需补充精确名称，不猜测。'),
 ('REQ AI 04 可复核','人工与自然语言入口调用同一业务服务；展示条件、来源流水、更新时间、查询编号。新会话可清除上下文。'),
 ('REQ VIS 01 看板','小时末道放行、工单状态、质量趋势及缺陷、物料齐套、库存安全线等有效视图；看板每30秒刷新，真实日期与班次筛选。'),
 ('REQ AUTO 01 自动化','完整报工且工序需检验时创建事件，提交后生成唯一待办；检验不合格保留待复检，合格关闭待办并推进进度。')],[3.4,13.2])
page(d,'二 权限和异常条件')
p(d,'库存与缺料开放给生产、车间、仓储、采购、业务管理员和财务岗位。质检与销售岗位可做单据追溯。操作员不开放该综合查询入口。所有查询仍须登录，服务端重新校验工具权限；会话只能由创建者继续使用。')
p(d,'规则变更由生产经理、车间负责人或业务管理员执行；质检员可处理或重试质检事件。采购员或业务管理员确认安全库存建议，确认本身不创建采购单。自动化页面每10秒刷新，人工可立即处理待执行队列。')
h(d,'三 成本和模型运行模式')
p(d,'默认采用明确标注的本地规则解析，无模型Token消耗。远程模式使用服务器环境变量设置地址、模型和密钥；密钥不进入前端和文档。每人每日最多20次远程调用，问题最多1200字，单次最大输出512 Token，接口等待最多10秒。失败调用也占次数，避免无限重试造成费用。实际输入输出Token来自接口usage，不估算金额。')
p(d,'接口超时、非成功状态、坏格式、非白名单字段或达到调用限额时返回明确失败，不自动以本地答案冒充远程模型成功。人工条件查询仍可使用。当前没有配置真实模型服务，已完成本地查询与接口协议模拟验收，真实服务连通性须配置后验证。')
h(d,'四 历史库存和业务边界')
p(d,'增量迁移只在首次执行时保存当前库存基线和流水水位。查询时间早于基线时，期初与期末显示无法核实，已有出入库只代表已保留流水，不能证明更早数据完整。禁止倒填正式基线以制造完整历史。')
p(d,'当前仓储业务采用整批领料、整批完工入库、整单发货。领料直接扣原料仓，不构成完整线边库存转移；分批领料、80/100部分发货以及ESP32采集未在本次完成范围内。库存图显示全部仓库当前余额，不跟随生产班次日期变化，页面已注明。')
h(d,'五 验收依据')
p(d,f'最终隔离库业务与查询验收 {accept["passed"]} 项通过；自动化故障注入 {fault["passed"]} 项通过；远程协议模拟 {remote["passed"]} 项通过。另有51项已有后端单元测试及前端类型检查与构建通过。真实模型联调和教师现场随机问题不在已通过项中。验收明细见同目录JSON和《验收与答辩手册》。')
d.save(OUT/'MouseMES软件需求规格补充说明.docx')

d=doc('MouseMES系统设计补充说明','第三阶段与第四阶段实现设计  版本 3.4  2026年10月23日')
p(d,'本说明修订原设计文档2.2.6、2.2.7及新增接口部分。代码采用Java 17、Spring Boot 3、JdbcTemplate和既有Vue 3页面，数据库保持MySQL与Redis部署方式。以下名称对应本次实际代码与增量数据表。')
h(d,'一 查询路径与职责')
p(d,'自然语言问题 → IntelligenceService解析与会话合并 → 条件白名单与角色校验 → BusinessQueryService业务查询 → 结构化结果与来源。人工表单直接进入同一个BusinessQueryService，不另设一套统计SQL。模型仅负责解析问题，数量计算始终在服务端完成。')
table(d,['组件','职责'],[
 ('BusinessQueryService','固定库存、缺料、追溯工具；参数绑定、日期验证、角色限制、基线核算、来源流水。'),
 ('IntelligenceService','本地规则或兼容Chat Completions接口；会话所有权、版本冲突、日调用额度、Token和耗时审计。'),
 ('AutomationService','报工事务中的事件落库；后台每2秒消费已提交事件；待办去重、规则开关、失败重试、检验结果回写。'),
 ('FoundationService','沿用任务报工与检验规则，进度同步处采集待检事件；过程检验完成时同步待办状态。'),
 ('CanonicalDashboardService','从正式mm_work_order、检验放行和库存台账统计；小时产量仅计末道工序放行，不累加所有工序报工。')],[5,11.6])
h(d,'二 前端入口')
p(d,'/intelligence提供自然语言、人工条件核对和本人用量记录；/automation展示规则、质检待办、执行日志和安全库存建议；/dashboard新增库存与安全线。待办检验链接附带workOrder与taskId，打开制造页面后定位工序并预选过程检验模板及来源。')
page(d,'三 数据结构与库存核算')
table(d,['表或约束','用途'],[
 ('mm_inventory_anchor','首次迁移的基线时刻与最大流水ID，只插入一次。'),
 ('mm_inventory_baseline','按物料、仓库保存基线数量，主键为二者组合。'),
 ('mm_ai_session','用户、筛选条件JSON和revision；更新以旧版本为条件，防止并发覆盖。'),
 ('mm_query_audit / mm_ai_budget','本人问题、条件、模式、状态、Token、耗时；用户加日期唯一额度，短事务原子预占调用次数。'),
 ('mm_auto_rule / mm_auto_event','规则开关；规则与task_id唯一，最多三次执行，记录状态、错误和处理时间。'),
 ('mm_qc_todo','task_id唯一、event_id唯一；WAITING、BLOCKED、DONE表示待检、待复检、已放行。'),
 ('mm_replenishment','每物料唯一确认记录；建议值实时从当前库存与安全库存计算。')],[6,10.6])
p(d,'可信时点余额 = 基线数量 + 水位之后、查询时点之前的入库 − 同范围出库。期间入出库按created_at左闭右开统计，用户输入的结束日期转换到次日零时。早于基线的期初与期末返回null；当前数量直接读取t_inventory。库存查询处于只读事务中，以一致快照读取基线、余额和流水。')
p(d,'原始库存流水按物料、仓库、库位和批次记录；查询汇总到物料与仓库，保留来源流水作为证据。一次最多200种物料；来源界面展示最新200条，超过限制会提示缩小条件。缺料比较全部未领料工单需求及安全库存与实际余额；指定仓库仅缩小可用量，需求仍为全部未领料工单，界面结果注明。')
h(d,'四 模型适配与边界')
p(d,'服务器变量为MES_AI_MODE、MES_AI_BASE_URL、MES_AI_MODEL、MES_AI_API_KEY。BASE_URL填写兼容接口的根地址，例如https://服务地址/v1，程序追加/chat/completions。远程请求采用JSON对象返回格式、512输出Token上限和10秒超时，只接收既定条件字段。接口不支持该协议时明确报错；不会发送业务结果表或数据库凭据给模型。')
page(d,'五 事务事件与自动化链')
p(d,'完整报工进入待质检状态时，在原报工事务内INSERT IGNORE写入REPORT_QC事件。事务回滚则事件同时回滚。后台按PENDING状态读取已提交事件，使用行锁与SKIP LOCKED防止多个工作线程重复消费。工序仍需检验且有有效过程模板时创建待办，再记执行成功。')
p(d,'事件与待办分别以规则加任务、任务和事件唯一约束去重。工序已放行或不需检验时跳过，规则停用时新事件记录SKIPPED。配置缺失或写入错误产生FAILED与可读原因；人工重试前确认规则启用且attempts小于3。第三次失败后保留日志交由人工处理，系统不修改检验事实以绕过失败。')
p(d,'检验不合格将待办置为BLOCKED，工序不能放行；复检合格在原检验事务内将待办置为DONE并保存检验ID与完成时刻。待办生成延迟取事件created_at到待办created_at的真实差值；该值不是人工检验耗时，也不能解释为产线效率提升比例。')
h(d,'六 实际接口契约')
table(d,['接口','用途与主要返回'],[
 ('POST /api/ai/query','question、可选sessionId；返回sessionId、traceId、mode、conditions、clarifications、result及Token。'),
 ('POST /api/analytics/query','kind及material、warehouse、product、type、from、to或document；返回同口径rows与sources。'),
 ('GET /api/ai/config /history','公开脱敏配置状态；history仅显示本人最近30条审计。'),
 ('GET /api/automation','rules、todos、events，日志最多200条。'),
 ('POST /api/automation/rule','enabled布尔值，由生产管理岗位控制。'),
 ('POST /api/automation/process','立即处理队列，用于演示或人工触发，不绕过规则校验。'),
 ('POST /api/automation/retry/{id}','失败或跳过事件重试，次数与角色受限。'),
 ('GET /api/automation/suggestions','实时安全库存缺口；/{id}/confirm确认，不创建采购单。')],[7,9.6])
p(d,'数据库脚本为04-database/stage34_intelligence.sql。迁移工具先完整备份，再执行两遍验证幂等，并核对原有业务表行数。正式迁移时先停止库存写入。回退应停止新服务并用备份恢复到独立恢复库核对；本次未执行任何覆盖式恢复。')
d.save(OUT/'MouseMES系统设计补充说明.docx')

d=doc('MouseMES第三四阶段验收与答辩手册','功能证据 演示顺序与部署说明  2026年10月23日')
p(d,'第三阶段已完成受限自然语言查询、人工对账、多轮澄清和模型适配；第四阶段已完成真实图表、报工自动质检待办和安全库存建议。本次交付以运行结果为依据，不保证教师最终评分。真实模型接口尚未配置，本地规则模式可直接使用。')
h(d,'一 测试结论与证据')
table(d,['验证','结果','范围'],[
 ('业务与查询验收',f'{accept["passed"]}项通过','真实订单到采购、领料、报工、检验、入库和发货；库存对账、多轮和权限。'),
 ('自动化故障注入',f'{fault["passed"]}项通过','缺失模板、失败记录、三次重试上限、恢复及唯一待办。'),
 ('模型协议模拟',f'{remote["passed"]}项通过','使用本地模拟服务，验证Token、非白名单、坏格式、503、超时和日限额。'),
 ('既有单元测试','51项通过','已有库存服务、MD5与头像测试；不是51项新增业务验收。'),
 ('前端验证','通过','Vue类型检查与生产构建；页面实测查询和自动化入口。'),
 ('真实模型联调','待配置','没有真实密钥，不声称已接入付费大模型或测得实际费用。')],[4,3,9.6])
h(d,'二 对应评分证据')
table(d,['评分方向','现场展示'],[
 ('智能查询准确性','时间库存、产品或类别缺料、工单追溯各选一题；切人工条件入口逐字段核对。'),
 ('结构化与多轮','展示表格、单位、日期与来源；多仓模糊条件追问，然后修改月份保留物料。'),
 ('可视化','小时末道产量、工单状态、齐套、质量趋势与缺陷、库存安全线；切日期、班次并核对数据。'),
 ('流程优化与自动化','报工事务提交后创建待办，日志显示唯一事件；不合格阻断、复检放行；演示失败后重试。')],[5,11.6])
page(d,'三 十五分钟演示脚本')
p(d,'优先使用5174隔离演示入口；页面黄色横幅标明测试数据，数据由真实业务接口生成。正式5173入口保持原有数据，不以验收测试单据冒充实际生产。先确认日期与隔离库工单计划日期一致。')
table(d,['时间','操作和讲解'],[
 ('0至2分钟','说明A100鼠标制造场景、角色、BOM与工艺路线；在生产与仓储闭环展示订单、采购和库存来源。'),
 ('2至5分钟','查看验收工单四道工序与三类检验，说明中间工序不增加成品完成量；展示领料、成品入库和发货流水。'),
 ('5至8分钟','智能查询输入“今天库存 物料：A100鼠标成品”，人工使用相同物料和日期核对入100、出100、余额0；再输入“改查上个月”，说明无法核实早于基线的余额。'),
 ('8至10分钟','输入“当前A100缺料”和“追溯 完整工单号”；展示报工、过程检验与28条库存流水。输入“查询库存 仓库：仓”触发歧义澄清。'),
 ('10至13分钟','打开规则与自动化，展示四条成功事件及已放行待办；若现场追加报工，等待自动出现待办，再以质检员执行检验；展示失败与三次重试验收记录。'),
 ('13至15分钟','展示至少四类图表、刷新时间和日期班次筛选；解释末道放行统计与库存当前值；结束前展示测试结论和模型配置边界。')],[3,13.6])
h(d,'四 流程优化如何对照')
p(d,'改进前：质检员在制造工单列表中逐单打开工序，查找待检任务，再进入检验。改进后：完整报工形成集中待办，质检员直接点击“执行检验”，工单、任务、模板与数量自动定位。现场对同一任务分别演示两条入口，核对是否减少查找和切换。')
p(d,'自动生成耗时在待办表逐行展示，取实际数据库时间差。不能将该耗时与人工检验时间比较，也没有采集班组人工耗时或漏检率样本，因此不填写“提升百分比”“节省多少分钟”等推测指标。安全库存建议另提供确认入口，避免自动采购带来误操作。')
page(d,'五 启动与复现')
p(d,'正式系统：http://127.0.0.1:5173/login。运行项目根目录下“启动MouseMES.cmd”；后台环境为Java 17、MySQL、Redis和Vue 3。与老师截图中的Java 8和Vue 2技术栈不同，本项目沿用已确认的MouseMES技术栈。')
p(d,'隔离演示：http://127.0.0.1:5174/login。运行08-environment/Start-Stage34-Demo.ps1，连接最近一次已验证的mouse_mes_stage34_test_时间戳库。该脚本不重新造单、不修改正式库，不可将演示库改名为正式库。')
p(d,'管理员演示工号109001，密码123456，拥有全部业务入口。生产工人103001，质检104001，仓管105001，采购106001，销售107001；其他角色请在基础配置用户中确认。真实凭据以现有本机用户表为准，本次不修改密码。')
h(d,'验收命令顺序')
p(d,'以下操作在D:\\competition\\MouseMES\\Mouse-Production-MES-System执行，使用已配置的Python与Java环境。第一次运行stage34_database.py --test创建全新隔离库；验收服务应指向该库并使用8082端口，不能复用已经产生验收单据的库重复执行全链创建脚本。')
for command in ['07-tools-barcodes/stage34_database.py --test','07-tools-barcodes/verify_stage34.py','07-tools-barcodes/verify_stage34_faults.py','07-tools-barcodes/verify_stage34_remote.py']:
 p(d,command)
p(d,'更完整的启动与迁移记录保存在.runtime目录，交付目录提供脱敏的验收JSON。正式备份位于.runtime/backups，以stage34-database.json中的backup路径为准。备份只用于恢复核对，不应直接提交含业务信息的完整备份给外部模型。')
h(d,'远程模型配置')
p(d,'设置MES_AI_MODE=remote、MES_AI_BASE_URL=兼容接口根地址、MES_AI_MODEL=服务商实际模型名、MES_AI_API_KEY=本机环境中的密钥，然后重启后端。密钥不要粘贴到文档、聊天或前端。配置后先看/ai/config显示已配置，再提交一次问题核对模式、Token与调用审计。切回本地可将MODE设为local并重启。')
page(d,'六 答辩问答和已知限制')
table(d,['可能的问题','依据事实回答'],[
 ('是否由模型计算库存','否。模型只解析条件，固定业务工具读取台账和流水，人工表单使用同一工具。'),
 ('当前是否使用大模型','默认是本地规则解析，页面明确标注；远程适配器已模拟测试，真实服务需要服务器密钥。'),
 ('历史余额为何为空','系统只在迁移时建立可信库存基线，早于基线的历史不能从当前库存倒猜。已有流水仍可展示并说明完整性限制。'),
 ('四次报工为何只有100只','成品完成量与小时产量按末道工序检验放行统计；中间工序报工不能重复计成品。'),
 ('重复点击会否重复待办','事件以规则和任务唯一，待办以任务和事件唯一；重复消费和故障恢复都已验收。'),
 ('自动化失败怎么办','记录失败原因，配置修复后由有权用户重试；最多三次，超过后人工处理，不能绕过检验放行。'),
 ('有没有自动采购','只有安全库存缺口建议与人工确认，仍由采购员创建采购单；不自行下采购订单。')],[5,11.6])
p(d,'未在本次新增：任意自由问答、模型自动写业务数据、设备硬件采集、完整线边库存流转、分批领料及部分发货。没有录制备用演示视频，现场应按脚本演示；需要备用录像时应由当前隔离环境实际录制，不使用静态假数据替代操作。')
p(d,'原需求与设计说明书中与实现不一致的规划段落，应按本目录两份补充说明替换后再合并提交。它们是明确的版本修订材料，本次保留原文档供比对，没有静默覆盖原稿。')
d.save(OUT/'MouseMES第三四阶段验收与答辩手册.docx')
for name in ['stage34-acceptance.json','stage34-fault-acceptance.json','stage34-remote-acceptance.json','stage34-database.json']:
 shutil.copyfile(ROOT/'.runtime'/name,OUT/name)
print(str(OUT))
